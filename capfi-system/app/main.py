"""
main.py — Backend FastAPI para consulta de precatórios TRF1
"""

import asyncio
import json
import os
import re
import uuid
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import FastAPI, UploadFile, WebSocket, WebSocketDisconnect, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from scraper import processar_lista, format_cpf, clean_cpf

# ─── Configuração ─────────────────────────────────────────────────────────────
BASE_DIR    = Path(__file__).parent
UPLOAD_DIR  = BASE_DIR / "uploads"
RESULT_DIR  = BASE_DIR / "results"
STATIC_DIR  = BASE_DIR / "static"

UPLOAD_DIR.mkdir(exist_ok=True)
RESULT_DIR.mkdir(exist_ok=True)

app = FastAPI(title="TRF1 Precatórios API", version="1.0.0")

# Jobs em memória: job_id -> {status, registros, resultados, progresso, ...}
jobs: Dict[str, dict] = {}

# ─── Estilização do Excel de saída ────────────────────────────────────────────

def gerar_excel_resultado(registros_entrada: list, resultados: list, caminho_saida: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precatórios TRF1"

    colunas = [
        "NOME", "CPF",
        "Nº Processo PRECAT", "Nº Processo Originário",
        "Precatório Expedido?", "Ano Orçamentário",
        "Data da Proposta", "Status"
    ]
    larguras = [40, 18, 30, 32, 22, 18, 20, 22]

    # ── Cabeçalho ──
    fill_header = PatternFill("solid", fgColor="1F4E79")
    font_header = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, col_name in enumerate(colunas, 1):
        c = ws.cell(row=1, column=col_idx, value=col_name)
        c.fill = fill_header
        c.font = font_header
        c.alignment = align_center
        c.border = border

    ws.row_dimensions[1].height = 35

    # Cores de status
    CORES_STATUS = {
        "Encontrado":    "E2EFDA",  # verde claro
        "Pago/Excluir":  "FFDDC1",  # laranja claro
        "Sem PRECAT":    "F2F2F2",  # cinza
        "Não encontrado": "F2F2F2", # cinza
        "Timeout":       "FFF2CC",  # amarelo
        "Erro":          "FFE0E0",  # vermelho claro
    }

    # ── Dados ──
    resultado_map = {r["cpf"]: r for r in resultados}

    for row_idx, reg in enumerate(registros_entrada, 2):
        cpf_fmt = format_cpf(reg["cpf"])
        res = resultado_map.get(cpf_fmt, {})

        status = res.get("status", "")
        cor_hex = next(
            (v for k, v in CORES_STATUS.items() if k in status),
            "FFFFFF"
        )
        fill_data = PatternFill("solid", fgColor=cor_hex)
        font_data = Font(name="Arial", size=10)

        valores = [
            reg["nome"],
            cpf_fmt,
            res.get("processo_precat", ""),
            res.get("processo_originario", ""),
            res.get("precatorio_expedido", "Não"),
            res.get("ano_orcamentario", ""),
            res.get("data_proposta", ""),
            status,
        ]

        for col_idx, valor in enumerate(valores, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=valor)
            c.font = font_data
            c.fill = fill_data
            c.border = border
            c.alignment = Alignment(
                horizontal="center" if col_idx > 2 else "left",
                vertical="center"
            )

    # ── Larguras ──
    for col_idx, width in enumerate(larguras, 1):
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width

    ws.freeze_panes = "A2"
    wb.save(caminho_saida)


# ─── Ler Excel de entrada ─────────────────────────────────────────────────────

def ler_excel(caminho: str) -> list[dict]:
    wb = openpyxl.load_workbook(caminho)
    ws = wb.active
    registros = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nome = row[0]
        cpf  = row[1]
        if nome and cpf:
            registros.append({"nome": str(nome).strip(), "cpf": str(cpf).strip()})
    return registros


# ─── Endpoints ────────────────────────────────────────────────────────────────

@app.post("/api/upload")
async def upload_arquivo(file: UploadFile):
    """Recebe o Excel com NOME e CPF e cria um job."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Envie um arquivo .xlsx ou .xls")

    job_id = str(uuid.uuid4())[:8]
    caminho = UPLOAD_DIR / f"{job_id}_{file.filename}"

    with open(caminho, "wb") as f:
        shutil.copyfileobj(file.file, f)

    registros = ler_excel(str(caminho))
    if not registros:
        os.remove(caminho)
        raise HTTPException(400, "Planilha vazia ou sem colunas NOME/CPF")

    checkpoint_path = str(RESULT_DIR / f"checkpoint_{job_id}.json")

    jobs[job_id] = {
        "status":          "aguardando",
        "arquivo":         str(caminho),
        "nome_arquivo":    file.filename,
        "registros":       registros,
        "resultados":      [],
        "checkpoint_path": checkpoint_path,
        "progresso": {
            "total":           len(registros),
            "processados":     0,
            "encontrados":     0,
            "pagos":           0,
            "nao_encontrados": 0,
            "erros":           0,
            "inicio":          None,
            "concluido":       False,
            "historico":       [],
        },
    }

    return {
        "job_id": job_id,
        "total":  len(registros),
        "nome_arquivo": file.filename,
    }


@app.post("/api/iniciar/{job_id}")
async def iniciar_processamento(job_id: str):
    """Inicia o processamento em background."""
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(404, "Job não encontrado")
    if job["status"] == "processando":
        raise HTTPException(400, "Já está em processamento")

    job["status"] = "processando"
    job["progresso"]["inicio"] = datetime.now().isoformat()

    asyncio.create_task(_processar_job(job_id))
    return {"status": "iniciado"}


async def _processar_job(job_id: str):
    job = jobs[job_id]
    prog = job["progresso"]

    async def callback(item: dict):
        job["resultados"].append(item)
        prog["processados"] += 1
        status = item["status"]
        if "Encontrado" in status:
            prog["encontrados"] += 1
        elif "Pago" in status:
            prog["pagos"] += 1
        elif "Não encontrado" in status or "Sem PRECAT" in status:
            prog["nao_encontrados"] += 1
        else:
            prog["erros"] += 1
        prog["historico"].append(item)

    try:
        await processar_lista(
            job["registros"],
            on_progresso=callback,
            checkpoint_path=job.get("checkpoint_path"),
        )
        job["status"] = "concluido"
        prog["concluido"] = True

        # Gerar Excel de resultado
        saida = str(RESULT_DIR / f"resultado_{job_id}.xlsx")
        gerar_excel_resultado(job["registros"], job["resultados"], saida)
        job["arquivo_resultado"] = saida

        # Remover checkpoint após conclusão bem-sucedida
        cp = job.get("checkpoint_path")
        if cp and os.path.exists(cp):
            os.remove(cp)

    except Exception as e:
        job["status"] = f"erro: {str(e)[:100]}"
        prog["concluido"] = True
        # Checkpoint é mantido para permitir retomada


@app.websocket("/ws/{job_id}")
async def websocket_progresso(ws: WebSocket, job_id: str):
    """WebSocket que transmite o progresso em tempo real."""
    await ws.accept()
    job = jobs.get(job_id)
    if not job:
        await ws.close(code=4004)
        return

    ultimo_enviado = -1
    try:
        while True:
            prog = job["progresso"]
            processados = prog["processados"]

            if processados > ultimo_enviado:
                await ws.send_text(json.dumps({
                    "tipo":            "progresso",
                    "status_job":      job["status"],
                    "total":           prog["total"],
                    "processados":     processados,
                    "encontrados":     prog["encontrados"],
                    "pagos":           prog["pagos"],
                    "nao_encontrados": prog["nao_encontrados"],
                    "erros":           prog["erros"],
                    "concluido":       prog["concluido"],
                    "historico":       prog["historico"],
                }))
                ultimo_enviado = processados

            if prog["concluido"]:
                # Envia mensagem final
                await ws.send_text(json.dumps({"tipo": "concluido"}))
                break

            await asyncio.sleep(1)

    except WebSocketDisconnect:
        pass


@app.get("/api/status/{job_id}")
async def status_job(job_id: str):
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(404, "Job não encontrado")
    return {
        "status":     job["status"],
        "progresso":  job["progresso"],
        "resultado_disponivel": "arquivo_resultado" in job,
    }


@app.get("/api/download/{job_id}")
async def download_resultado(job_id: str):
    job = jobs.get(job_id)
    if not job:
        raise HTTPException(404, "Job não encontrado")
    if "arquivo_resultado" not in job:
        raise HTTPException(400, "Resultado ainda não disponível")

    nome_saida = f"Resultado_Precatorios_{job_id}.xlsx"
    return FileResponse(
        path=job["arquivo_resultado"],
        filename=nome_saida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ─── Servir frontend ──────────────────────────────────────────────────────────
app.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="static")
