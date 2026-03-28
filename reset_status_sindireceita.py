"""
Reseta PROCESSO PRECAT, ORÇAMENTO e STATUS_CONSULTA do SINDIRECEITA_COMPLETO.xlsx
para que o script consultar_sindireceita_trf1.py reprocesse todas as 445 linhas do zero.
"""
from openpyxl import load_workbook
from pathlib import Path

EXCEL_PATH = "SINDIRECEITA_COMPLETO.xlsx"
COL_PRECAT  = 4
COL_ORC     = 5
COL_STATUS  = 12

excel = Path(EXCEL_PATH)
if not excel.exists():
    print(f"❌ Arquivo não encontrado: {EXCEL_PATH}")
    print(f"   Pasta atual: {Path.cwd()}")
    exit(1)

wb = load_workbook(EXCEL_PATH)
ws = wb.active

count = 0
for row in range(2, ws.max_row + 1):
    ws.cell(row, COL_PRECAT).value = None
    ws.cell(row, COL_ORC).value    = None
    ws.cell(row, COL_STATUS).value = None
    count += 1

wb.save(EXCEL_PATH)
print(f"✅ {count} linhas resetadas (PRECAT, ORÇAMENTO e STATUS limpos).")
print(f"   Agora rode: caffeinate -i python3 consultar_sindireceita_trf1.py")
