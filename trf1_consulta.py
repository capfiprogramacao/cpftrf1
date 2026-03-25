"""
TRF1 - Consulta de Precatórios por CPF
Automação com Playwright + atualização de planilha Excel
"""

import asyncio
import re
import json
import os
from datetime import datetime
from pathlib import Path

from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout
from openpyxl import load_workbook

# ─── Configurações ────────────────────────────────────────────────────────────
BASE_DIR    = Path(__file__).parent
EXCEL_PATH  = str(BASE_DIR / "mnt/outputs/Aposentados Sindireceita Teste.xlsx")
PROGRESS_FILE = str(BASE_DIR / "progress.json")
URL_CPF_BUSCA = "https://processual.trf1.jus.br/consultaProcessual/cpfCnpjParte.php?secao=TRF1"

# Palavras-chave que indicam precatório encerrado (pagamento/levantamento)
KEYWORDS_PAGO = [
    "depósito", "deposito", "saque", "pagamento", "pago",
    "levantamento", "alvará", "alvar", "requisição de pequeno valor",
    "rpv levantado", "quitação", "quitado", "crédito levantado",
    "transferência", "ordem de pagamento expedida"
]

# Palavra-chave para proposta orçamentária
KEYWORD_PROPOSTA = "proposta orçamentária enviada ao cjf"

# ─── Utilitários ──────────────────────────────────────────────────────────────

def clean_cpf(cpf: str) -> str:
    return re.sub(r'\D', '', str(cpf).strip())

def format_cpf(cpf: str) -> str:
    digits = clean_cpf(cpf)
    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    return cpf

def clean_processo_originario(texto: str) -> str:
    if texto and '/' in texto:
        return texto.split('/')[0].strip()
    return texto.strip() if texto else ""

def save_progress(data: dict):
    with open(PROGRESS_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def load_excel_data(retomar=True):
    """
    Carrega os registros da planilha.
    Se retomar=True, pula linhas que já foram processadas com sucesso,
    permitindo continuar de onde parou em caso de queda de energia/internet.
    Linhas com status "Erro" ou "Timeout" são reprocessadas automaticamente.
    """
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active

    # Garantir que os cabeçalhos existam
    if ws.cell(row=1, column=1).value is None:
        ws.cell(row=1, column=1).value = "NOME"
        ws.cell(row=1, column=2).value = "CPF"
        ws.cell(row=1, column=3).value = "Nº Processo PRECAT"
        ws.cell(row=1, column=4).value = "Nº Processo Originário"
        ws.cell(row=1, column=5).value = "Precatório Expedido?"
        ws.cell(row=1, column=6).value = "Ano Orçamentário"
        ws.cell(row=1, column=7).value = "Data da Proposta"
        ws.cell(row=1, column=8).value = "Status"
        wb.save(EXCEL_PATH)

    registros = []
    ja_processados = 0
    STATUS_REPROCESSAR = {"Erro", "Timeout", ""}  # Status que devem ser reprocessados

    for row_idx in range(2, ws.max_row + 1):
        nome   = ws.cell(row=row_idx, column=1).value
        cpf    = ws.cell(row=row_idx, column=2).value
        status = ws.cell(row=row_idx, column=8).value

        if not nome or not cpf:
            continue

        status_atual = str(status).strip() if status else ""

        # Se retomar está ativo e a linha já tem status válido, pular
        if retomar and status_atual and status_atual not in STATUS_REPROCESSAR:
            ja_processados += 1
            continue

        registros.append((row_idx, str(nome).strip(), str(cpf).strip()))

    if retomar and ja_processados > 0:
        print(f"\n  ♻️  Retomando processamento anterior...")
        print(f"  ✅ Já processados: {ja_processados} registros (serão ignorados)")
        print(f"  🔄 Pendentes: {len(registros)} registros\n")

    return registros

def update_excel_row(row_idx: int, dados: dict):
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    ws.cell(row=row_idx, column=3).value = dados.get("processo_precat", "")
    ws.cell(row=row_idx, column=4).value = dados.get("processo_originario", "")
    ws.cell(row=row_idx, column=5).value = dados.get("precatorio_expedido", "")
    ws.cell(row=row_idx, column=6).value = dados.get("ano_orcamentario", "")
    ws.cell(row=row_idx, column=7).value = dados.get("data_proposta", "")
    ws.cell(row=row_idx, column=8).value = dados.get("status", "")
    wb.save(EXCEL_PATH)

# ─── Automação Principal ──────────────────────────────────────────────────────

async def consultar_cpf(page, cpf: str, nome: str, primeira_vez: bool) -> dict:
    resultado = {
        "processo_precat": "",
        "processo_originario": "",
        "precatorio_expedido": "Não",
        "ano_orcamentario": "",
        "data_proposta": "",
        "status": "Não encontrado"
    }

    try:
        cpf_digits = clean_cpf(cpf)

        # ── ETAPA 1: Navegar para busca por CPF (apenas primeira vez) ──────────
        if primeira_vez:
            await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(3)

        # ── ETAPA 2: Preencher formulário ─────────────────────────────────────
        # Marcar "Mostrar os baixados"
        checkbox = page.locator('input[name="mostrarBaixados"]')
        if await checkbox.count() > 0 and not await checkbox.is_checked():
            await checkbox.check()

        # Limpar e preencher CPF (campo correto: cpf_cnpj)
        cpf_input = page.locator('input[name="cpf_cnpj"]')
        await cpf_input.click()
        await cpf_input.evaluate('el => el.value = ""')
        await cpf_input.fill(cpf_digits)

        # ── ETAPA 3: Pesquisar ────────────────────────────────────────────────
        await page.locator('input#enviar').click()
        await page.wait_for_load_state('networkidle', timeout=30000)
        await asyncio.sleep(2)

        # Verificar se encontrou partidas
        body_text = await page.evaluate('document.body.innerText')
        if 'partes encontradas' not in body_text.lower() and 'nome da parte' not in body_text.lower():
            resultado["status"] = "Não encontrado"
            # Voltar para o formulário para o próximo CPF
            await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        # ── ETAPA 4: Clicar na parte para expandir a lista de processos ───────
        parte_links = page.locator('table a')
        count = await parte_links.count()
        if count == 0:
            resultado["status"] = "Não encontrado"
            await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        await parte_links.first.click()
        await page.wait_for_load_state('networkidle', timeout=20000)
        await asyncio.sleep(2)

        # ── ETAPA 5: Identificar processo PRC na lista ────────────────────────
        # Ler a tabela de processos
        rows_data = await page.evaluate('''
            () => {
                const rows = Array.from(document.querySelectorAll("table tr"));
                return rows.map(row => {
                    const cells = Array.from(row.querySelectorAll("td"));
                    const link = row.querySelector("a");
                    return {
                        col1: cells[0] ? cells[0].innerText.trim() : "",
                        col2: cells[1] ? cells[1].innerText.trim() : "",
                        href: link ? link.href : null,
                        text: link ? link.innerText.trim() : ""
                    };
                }).filter(r => r.href && r.col1.length > 0);
            }
        ''')

        # Pegar o ÚLTIMO processo PRC (mais recente) — pode haver múltiplos
        precat_rows = [r for r in rows_data if '(PRC)' in r['col1'] or '(PREC)' in r['col1']]
        precat_row = precat_rows[-1] if precat_rows else None

        if not precat_row:
            resultado["status"] = "Sem PRECAT"
            await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        # Extrair dados da lista
        resultado["processo_precat"] = precat_row['col1'].strip()
        resultado["processo_originario"] = clean_processo_originario(precat_row['col2'])
        resultado["precatorio_expedido"] = "Sim"

        # ── ETAPA 6: Abrir processo PRC clicando no link (último da lista) ───
        todos_prc = page.locator('table a', has_text='PRC')
        total_prc = await todos_prc.count()
        prc_link = todos_prc.nth(total_prc - 1) if total_prc > 0 else todos_prc.first
        await prc_link.click()
        await page.wait_for_load_state('networkidle', timeout=30000)
        await asyncio.sleep(2)

        # ── ETAPA 7: Clicar na aba Movimentação ──────────────────────────────
        mov_link = page.locator('a', has_text='Movimentação')
        if await mov_link.count() == 0:
            resultado["status"] = "Encontrado - Sem aba Movimentação"
            await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        await mov_link.click()
        await page.wait_for_load_state('networkidle', timeout=20000)
        await asyncio.sleep(2)

        # ── ETAPA 8: Extrair dados da movimentação ────────────────────────────
        mov_text = await page.evaluate('document.body.innerText')
        mov_lower = mov_text.lower()

        # Buscar proposta orçamentária PRIMEIRO (antes de checar pagamento)
        mov_rows = await page.evaluate('''
            () => {
                const rows = Array.from(document.querySelectorAll("table tr"));
                return rows.map(row => {
                    const cells = Array.from(row.querySelectorAll("td"));
                    return cells.map(c => c.innerText.trim());
                }).filter(r => r.length >= 3);
            }
        ''')

        proposta_encontrada = False
        for row in mov_rows:
            row_text = ' '.join(row).lower()
            if 'proposta orçamentária' in row_text and 'cjf' in row_text:
                proposta_encontrada = True

                # Complemento está na última coluna (ex: "2027,data 13/02/2026")
                complemento = row[-1] if row else ''

                # Extrair ano orçamentário: 4 dígitos no início do complemento
                match_ano = re.match(r'(\d{4})', complemento.strip())
                if match_ano:
                    resultado["ano_orcamentario"] = match_ano.group(1)
                else:
                    anos = re.findall(r'\b(20\d{2})\b', complemento)
                    if anos:
                        resultado["ano_orcamentario"] = anos[0]

                # Extrair data (dd/mm/yyyy) após "data " no complemento
                match_data = re.search(r'data\s+(\d{2}/\d{2}/\d{4})', complemento, re.I)
                if match_data:
                    resultado["data_proposta"] = match_data.group(1)
                else:
                    data_col = row[0] if row else ''
                    match_data2 = re.search(r'(\d{2}/\d{2}/\d{4})', data_col)
                    if match_data2:
                        resultado["data_proposta"] = match_data2.group(1)
                break

        # Verificar palavras-chave de pagamento/encerramento APÓS extrair dados
        status_pago = any(kw in mov_lower for kw in KEYWORDS_PAGO)

        if status_pago:
            resultado["status"] = "Pago/Excluir"
        elif proposta_encontrada:
            resultado["status"] = "Encontrado"
        else:
            resultado["status"] = "Encontrado - Sem proposta orçamentária"

        # Voltar para o formulário para o próximo CPF
        await page.goto(URL_CPF_BUSCA, timeout=30000, wait_until='networkidle')
        await asyncio.sleep(2)

    except PlaywrightTimeout:
        resultado["status"] = "Timeout"
        print(f"  ⚠️  Timeout para CPF {cpf}")
        try:
            await page.goto(URL_CPF_BUSCA, timeout=15000, wait_until='domcontentloaded')
            await asyncio.sleep(2)
        except Exception:
            pass

    except Exception as e:
        resultado["status"] = "Erro"
        print(f"  ❌ Erro: {e}")
        try:
            await page.goto(URL_CPF_BUSCA, timeout=15000, wait_until='domcontentloaded')
            await asyncio.sleep(2)
        except Exception:
            pass

    return resultado


async def processar_planilha():
    registros = load_excel_data(retomar=True)

    # Contar total real da planilha para exibir progresso correto
    wb_total = load_workbook(EXCEL_PATH)
    ws_total = wb_total.active
    total_planilha = sum(
        1 for r in range(2, ws_total.max_row + 1)
        if ws_total.cell(row=r, column=1).value and ws_total.cell(row=r, column=2).value
    )
    ja_feitos = total_planilha - len(registros)
    total = len(registros)

    if total == 0:
        print(f"\n{'='*60}")
        print(f"  ✅ Todos os {total_planilha} registros já foram processados!")
        print(f"  Nenhuma ação necessária.")
        print(f"{'='*60}\n")
        return

    print(f"\n{'='*60}")
    print(f"  TRF1 — Consulta de Precatórios por CPF")
    print(f"  Total na planilha   : {total_planilha}")
    print(f"  Já processados      : {ja_feitos}")
    print(f"  A processar agora   : {total}")
    print(f"  Início: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    print(f"{'='*60}\n")

    progresso = {
        "total": total_planilha, "processados": ja_feitos, "encontrados": 0,
        "nao_encontrados": 0, "pagos": 0, "erros": 0,
        "inicio": datetime.now().isoformat(),
        "ultimo_update": datetime.now().isoformat(),
        "concluido": False, "registros": []
    }
    save_progress(progresso)

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage", "--disable-blink-features=AutomationControlled"]
        )
        context = await browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 800},
            locale="pt-BR"
        )
        page = await context.new_page()
        await page.add_init_script('Object.defineProperty(navigator, "webdriver", {get: () => undefined})')

        primeira_vez = True
        for idx, (row_idx, nome, cpf) in enumerate(registros):
            print(f"[{idx+1:03d}/{total}] {nome} | CPF: {format_cpf(cpf)}")

            dados = await consultar_cpf(page, cpf, nome, primeira_vez)
            primeira_vez = False

            update_excel_row(row_idx, dados)

            progresso["processados"] += 1
            progresso["ultimo_update"] = datetime.now().isoformat()
            status = dados["status"]

            if "Encontrado" in status:
                progresso["encontrados"] += 1
            elif "Pago" in status:
                progresso["pagos"] += 1
            elif "Não encontrado" in status or "Sem PRECAT" in status:
                progresso["nao_encontrados"] += 1
            else:
                progresso["erros"] += 1

            progresso["registros"].append({
                "nome": nome, "cpf": format_cpf(cpf), "status": status,
                "processo_precat": dados.get("processo_precat", ""),
                "ano_orcamentario": dados.get("ano_orcamentario", ""),
                "hora": datetime.now().strftime("%H:%M:%S")
            })
            save_progress(progresso)

            emoji = {"Encontrado": "✅", "Pago/Excluir": "🔴", "Não encontrado": "⚪",
                     "Sem PRECAT": "⚪", "Erro": "❌", "Timeout": "⏱️"}
            e = next((v for k, v in emoji.items() if k in status), "❓")
            print(f"       {e} {status} | PRECAT: {dados.get('processo_precat', '-')} | Ano: {dados.get('ano_orcamentario', '-')}\n")

            # Pausa entre consultas
            await asyncio.sleep(1.5)

        await browser.close()

    print(f"\n{'='*60}")
    print(f"  ✅ Processamento concluído!")
    print(f"  Encontrados com PRECAT : {progresso['encontrados']}")
    print(f"  Pagos / Excluir        : {progresso['pagos']}")
    print(f"  Não encontrados        : {progresso['nao_encontrados']}")
    print(f"  Erros                  : {progresso['erros']}")
    print(f"  Término: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    print(f"{'='*60}\n")
    print(f"📄 Planilha atualizada: {EXCEL_PATH}")

    progresso["concluido"] = True
    save_progress(progresso)


if __name__ == "__main__":
    asyncio.run(processar_planilha())
