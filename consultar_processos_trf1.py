"""
CAPFI — Consulta TRF1 por Número de Processo
Atualiza PROCESSO PRECAT (col 4) e ANO ORÇAMENTO (col 5)

Como usar:
  1. Coloque este script na mesma pasta que SINDIRECEITA_COMPLETO.xlsx
  2. Execute: caffeinate -i python3 consultar_processos_trf1.py
"""

import asyncio
import re
import sys
from datetime import datetime
from pathlib import Path
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout
from openpyxl import load_workbook

# ─── CONFIGURAÇÃO ──────────────────────────────────────────────
EXCEL_PATH  = "SINDIRECEITA_COMPLETO.xlsx"
URL         = "https://processual.trf1.jus.br/consultaProcessual/numeroProcesso.php?secao=TRF1"
NUM_WORKERS = 5
# ───────────────────────────────────────────────────────────────

STATUS_SKIP = {"OK", "Sem PRECAT", "Não encontrado", "Erro permanente"}

def limpar_processo(proc):
    """Remove caracteres extras e formata o número do processo."""
    return re.sub(r'[^\d\-\.]', '', str(proc).strip())

excel_lock = asyncio.Lock()

async def salvar_resultado(row, precat, orcamento, status):
    async with excel_lock:
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        if precat:
            ws.cell(row, 4).value = precat
        if orcamento:
            ws.cell(row, 5).value = orcamento
        ws.cell(row, 12).value = status  # coluna extra de controle
        wb.save(EXCEL_PATH)

async def consultar_processo(page, numero_processo, primeira_vez):
    resultado = {"precat": None, "orcamento": None, "status": "Não encontrado"}

    try:
        if primeira_vez:
            await page.goto(URL, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(3)

        # Preenche o campo de número do processo
        inp = page.locator('input#proc')
        await inp.click()
        await inp.evaluate('el => el.value = ""')
        await inp.fill(limpar_processo(numero_processo))

        # Clica em pesquisar
        await page.locator('input#enviar').click()
        await page.wait_for_load_state('networkidle', timeout=30000)
        await asyncio.sleep(2)

        body = await page.evaluate('document.body.innerText')
        body_low = body.lower()

        if 'nenhum processo encontrado' in body_low or 'não foram encontrados' in body_low:
            await page.goto(URL, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        # Extrai processos da tabela
        rows_data = await page.evaluate('''() => {
            return Array.from(document.querySelectorAll("table tr")).map(row => {
                const cols = Array.from(row.querySelectorAll("td"));
                const link = row.querySelector("a");
                return {
                    col1: cols[0] ? cols[0].innerText.trim() : "",
                    col2: cols[1] ? cols[1].innerText.trim() : "",
                    href: link ? link.href : null
                };
            }).filter(r => r.href && r.col1.length > 0);
        }''')

        # Procura processo PRECAT (PRC ou PREC)
        precat_rows = [r for r in rows_data if '(PRC)' in r['col1'] or '(PREC)' in r['col1']]

        if not precat_rows:
            resultado["status"] = "Sem PRECAT"
            await page.goto(URL, timeout=30000, wait_until='networkidle')
            await asyncio.sleep(2)
            return resultado

        pr = precat_rows[-1]
        resultado["precat"] = pr['col1'].strip()

        # Abre o processo PRECAT para pegar o ano orçamentário
        todos = page.locator('table a', has_text='PRC')
        n = await todos.count()
        await (todos.nth(n - 1) if n > 0 else todos.first).click()
        await page.wait_for_load_state('networkidle', timeout=30000)
        await asyncio.sleep(2)

        # Verifica aba Movimentação para encontrar o ano orçamentário
        mov = page.locator('a', has_text='Movimentação')
        if await mov.count() > 0:
            await mov.click()
            await page.wait_for_load_state('networkidle', timeout=20000)
            await asyncio.sleep(2)

            mr = await page.evaluate('''() =>
                Array.from(document.querySelectorAll("table tr"))
                .map(row => Array.from(row.querySelectorAll("td"))
                .map(c => c.innerText.trim()))
                .filter(r => r.length >= 3)''')

            for r in mr:
                joined = ' '.join(r).lower()
                if 'proposta orçamentária' in joined and 'cjf' in joined:
                    anos = re.findall(r'\b(20\d{2})\b', r[-1] if r else '')
                    if anos:
                        resultado["orcamento"] = anos[0]
                    break

        resultado["status"] = "OK"
        await page.goto(URL, timeout=30000, wait_until='networkidle')
        await asyncio.sleep(2)

    except PlaywrightTimeout:
        resultado["status"] = "Timeout"
        print(f"    ⚠️  Timeout: {numero_processo}", flush=True)
        try:
            await page.goto(URL, timeout=15000, wait_until='domcontentloaded')
            await asyncio.sleep(2)
        except: pass

    except Exception as e:
        resultado["status"] = "Erro"
        print(f"    ❌ Erro em {numero_processo}: {e}", flush=True)
        try:
            await page.goto(URL, timeout=15000, wait_until='domcontentloaded')
            await asyncio.sleep(2)
        except: pass

    return resultado

async def worker(worker_id, fila, lock, playwright, contagem):
    browser = await playwright.chromium.launch(
        headless=True,
        args=[
            "--no-sandbox",
            "--disable-dev-shm-usage",
            "--disable-blink-features=AutomationControlled"
        ]
    )
    ctx = await browser.new_context(
        user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        viewport={"width": 1280, "height": 800},
        locale="pt-BR"
    )
    page = await ctx.new_page()
    await page.add_init_script('Object.defineProperty(navigator, "webdriver", {get: () => undefined})')

    primeira_vez = True
    try:
        while True:
            try:
                row, nome, processo = fila.get_nowait()
            except Exception:
                break

            resultado = await consultar_processo(page, processo, primeira_vez)
            primeira_vez = False

            await salvar_resultado(row, resultado["precat"], resultado["orcamento"], resultado["status"])

            async with lock:
                contagem[0] += 1
                emoji = "✅" if resultado["status"] == "OK" else \
                        "⚪" if resultado["status"] in ("Sem PRECAT", "Não encontrado") else "❌"
                precat_str = resultado["precat"] or "-"
                orc_str    = resultado["orcamento"] or "-"
                print(f"[W{worker_id}][{contagem[0]:04d}] {nome[:30]:<30} | PRECAT: {precat_str[:30]} | Orç: {orc_str} {emoji}", flush=True)

            fila.task_done()
            await asyncio.sleep(1)
    finally:
        await browser.close()

async def main():
    excel = Path(EXCEL_PATH)
    if not excel.exists():
        print(f"❌ Arquivo não encontrado: {EXCEL_PATH}")
        print(f"   Pasta atual: {Path.cwd()}")
        sys.exit(1)

    print(f"📂 Arquivo: {excel.resolve()}")

    # Adiciona coluna de controle se necessário
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    if ws.cell(1, 12).value is None:
        ws.cell(1, 12).value = "STATUS_CONSULTA"
        wb.save(EXCEL_PATH)

    # Carrega pendentes (sem status ou com Erro/Timeout)
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    pendentes = []
    for row in range(2, ws.max_row + 1):
        nome     = ws.cell(row, 1).value
        processo = ws.cell(row, 6).value  # PROCESSO ANALISADO
        status   = str(ws.cell(row, 12).value or "").strip()
        if not processo: continue
        if status in STATUS_SKIP: continue
        pendentes.append((row, str(nome or "").strip(), str(processo).strip()))

    total  = ws.max_row - 1
    feitos = total - len(pendentes)

    print(f"\n{'='*65}")
    print(f"  CAPFI — TRF1 | Consulta por Processo")
    print(f"  Workers  : {NUM_WORKERS} navegadores em paralelo")
    print(f"  Total    : {total} | Feitos: {feitos} | Pendentes: {len(pendentes)}")
    print(f"  Início   : {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    print(f"{'='*65}\n")

    if not pendentes:
        print("✅ Todos os processos já foram consultados!")
        return

    fila = asyncio.Queue()
    for item in pendentes:
        fila.put_nowait(item)

    lock     = asyncio.Lock()
    contagem = [0]

    async with async_playwright() as p:
        await asyncio.gather(*[
            worker(i + 1, fila, lock, p, contagem)
            for i in range(NUM_WORKERS)
        ])

    print(f"\n🏁 Concluído! {contagem[0]} processos consultados.")
    print(f"📁 Salvo em: {excel.resolve()}")

if __name__ == "__main__":
    asyncio.run(main())
